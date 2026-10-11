"""Excel spreadsheet exporter for conversations and markdown tables using openpyxl."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.export.business_logic.base_exporter import BaseExporter
from app.export.business_logic.markdown_blocks import parse_markdown_blocks
from app.export.export_options import ExportOptions
from app.export.models import ExportConversation

COLOR_PALETTES: Dict[str, Dict[str, str]] = {
    "vibrant": {
        "title_fill": "4F46E5",
        "title_text": "FFFFFF",
        "header_fill": "7C3AED",
        "header_text": "FFFFFF",
        "alt_row": "F5F3FF",
        "total_fill": "DDD6FE",
        "total_text": "4C1D95",
        "border": "C7D2FE",
    },
    "colorful": {
        "title_fill": "7C3AED",
        "title_text": "FFFFFF",
        "header_fill": "2563EB",
        "header_text": "FFFFFF",
        "alt_row": "EFF6FF",
        "total_fill": "BFDBFE",
        "total_text": "1E40AF",
        "border": "93C5FD",
    },
    "emerald": {
        "title_fill": "065F46",
        "title_text": "FFFFFF",
        "header_fill": "059669",
        "header_text": "FFFFFF",
        "alt_row": "F0FDF4",
        "total_fill": "A7F3D0",
        "total_text": "064E3B",
        "border": "6EE7B7",
    },
    "sunset": {
        "title_fill": "9A3412",
        "title_text": "FFFFFF",
        "header_fill": "E11D48",
        "header_text": "FFFFFF",
        "alt_row": "FFF7ED",
        "total_fill": "FDE68A",
        "total_text": "78350F",
        "border": "FDBA74",
    },
    "cyber": {
        "title_fill": "0F172A",
        "title_text": "06B6D4",
        "header_fill": "1E293B",
        "header_text": "38BDF8",
        "alt_row": "F8FAFC",
        "total_fill": "E0F2FE",
        "total_text": "0369A1",
        "border": "7DD3FC",
    },
    "executive": {
        "title_fill": "0F172A",
        "title_text": "FFFFFF",
        "header_fill": "1E293B",
        "header_text": "FFFFFF",
        "alt_row": "F8FAFC",
        "total_fill": "E2E8F0",
        "total_text": "0F172A",
        "border": "CBD5E1",
    },
}


def resolve_palette(theme: str = "") -> Dict[str, str]:
    clean_theme = (theme or "").strip().lower()
    if clean_theme in COLOR_PALETTES:
        return COLOR_PALETTES[clean_theme]
    if "color" in clean_theme or "rainbow" in clean_theme or "multi" in clean_theme:
        return COLOR_PALETTES["vibrant"]
    if "green" in clean_theme or "mint" in clean_theme or "tree" in clean_theme:
        return COLOR_PALETTES["emerald"]
    if "orange" in clean_theme or "warm" in clean_theme or "red" in clean_theme:
        return COLOR_PALETTES["sunset"]
    if "dark" in clean_theme or "blue" in clean_theme:
        return COLOR_PALETTES["cyber"]
    return COLOR_PALETTES["executive"]


class ExcelExporter(BaseExporter):
    """Render conversation tables and structured content into a styled Excel workbook."""

    format_id = "xlsx"
    display_name = "Excel spreadsheet"
    file_extension = ".xlsx"

    def export(self, conversation: ExportConversation, destination: Path, options: ExportOptions) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"
        ws.views.sheetView[0].showGridLines = True

        palette = resolve_palette(getattr(options, "theme", "executive"))
        current_row = 1

        # Title Block
        if not options.content_only:
            ws.cell(row=current_row, column=1, value=conversation.title)
            title_cell = ws.cell(row=current_row, column=1)
            title_cell.font = Font(name="Segoe UI", size=16, bold=True, color=palette["title_text"])
            title_cell.fill = PatternFill(start_color=palette["title_fill"], end_color=palette["title_fill"], fill_type="solid")
            title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[current_row].height = 36
            current_row += 2

        if options.include_metadata and not options.content_only:
            meta_text = f"ID: {conversation.id} | Created: {conversation.created_at} | Updated: {conversation.updated_at}"
            cell = ws.cell(row=current_row, column=1, value=meta_text)
            cell.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
            current_row += 2

        table_count = 0
        for message in conversation.messages:
            blocks = parse_markdown_blocks(message.content)
            for block in blocks:
                if block.kind == "table" and block.rows:
                    table_count += 1
                    current_row = self._write_styled_table(ws, block.rows, current_row, palette=palette, title=f"Table {table_count}")
                    current_row += 2

        # If no markdown tables were found, render all content into formatted rows
        if table_count == 0:
            for message in conversation.messages:
                role_str = str(getattr(message.role, "value", message.role)).capitalize()
                cell = ws.cell(row=current_row, column=1, value=role_str)
                cell.font = Font(name="Segoe UI", size=12, bold=True, color=palette["title_fill"])
                current_row += 1
                for line in message.content.splitlines():
                    if line.strip():
                        c = ws.cell(row=current_row, column=1, value=line.strip())
                        c.font = Font(name="Segoe UI", size=10, color="334155")
                        current_row += 1
                current_row += 1

        self._auto_fit_columns(ws)
        wb.save(str(destination))
        return destination

    @classmethod
    def create_excel_from_markdown(cls, content: str, destination: Path, title: str = "Smart Data Report", theme: str = "vibrant") -> Path:
        """Utility method to directly generate a colorful Excel workbook from Markdown text containing tables."""
        destination.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.views.sheetView[0].showGridLines = True

        palette = resolve_palette(theme)
        current_row = 1

        # Document Header
        ws.cell(row=current_row, column=1, value=title)
        title_cell = ws.cell(row=current_row, column=1)
        title_cell.font = Font(name="Segoe UI", size=15, bold=True, color=palette["title_text"])
        title_cell.fill = PatternFill(start_color=palette["title_fill"], end_color=palette["title_fill"], fill_type="solid")
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_row].height = 34
        current_row += 2

        blocks = parse_markdown_blocks(content)
        table_count = 0

        for block in blocks:
            if block.kind == "table" and block.rows:
                table_count += 1
                current_row = cls._write_styled_table(ws, block.rows, current_row, palette=palette)
                current_row += 2
            elif block.kind == "paragraph":
                # Paragraph text as section callouts or headings
                for line in block.text.splitlines():
                    clean = line.strip()
                    if clean.startswith("#"):
                        heading_text = clean.lstrip("#").strip()
                        c = ws.cell(row=current_row, column=1, value=heading_text)
                        c.font = Font(name="Segoe UI", size=12, bold=True, color=palette["title_fill"])
                        current_row += 1
                    elif clean:
                        c = ws.cell(row=current_row, column=1, value=clean)
                        c.font = Font(name="Segoe UI", size=10, italic=False, color="475569")
                        current_row += 1

        # Fallback if no table blocks detected
        if table_count == 0:
            lines = [l.strip() for l in content.splitlines() if l.strip()]
            for line in lines:
                c = ws.cell(row=current_row, column=1, value=line)
                c.font = Font(name="Segoe UI", size=10, color="334155")
                current_row += 1

        cls._auto_fit_columns(ws)
        wb.save(str(destination))
        return destination

    @classmethod
    def _write_styled_table(
        cls,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        rows: Tuple[Tuple[str, ...], ...],
        start_row: int,
        palette: Dict[str, str] = None,
        title: str = "",
    ) -> int:
        if not rows:
            return start_row

        if palette is None:
            palette = COLOR_PALETTES["executive"]

        header_fill = PatternFill(start_color=palette["header_fill"], end_color=palette["header_fill"], fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color=palette["header_text"])

        row_fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        row_fill_odd = PatternFill(start_color=palette["alt_row"], end_color=palette["alt_row"], fill_type="solid")
        data_font = Font(name="Segoe UI", size=10, color="334155")

        thin_side = Side(border_style="thin", color=palette["border"])
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        total_bottom_border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=Side(border_style="double", color=palette["title_fill"]),
        )

        align_center = Alignment(horizontal="center", vertical="center")

        current_r = start_row

        # Write Header Row
        headers = rows[0]
        ws.row_dimensions[current_r].height = 24
        for c_idx, val in enumerate(headers, start=1):
            cell = ws.cell(row=current_r, column=c_idx, value=val.strip())
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = align_center

        current_r += 1

        # Write Data Rows
        for r_idx, row in enumerate(rows[1:], start=0):
            ws.row_dimensions[current_r].height = 20
            fill = row_fill_odd if r_idx % 2 == 1 else row_fill_even
            is_total_row = any(str(v).strip().lower() in ("total", "sum", "average", "grand total") for v in row)

            for c_idx, val in enumerate(row, start=1):
                clean_val = val.strip()
                cell = ws.cell(row=current_r, column=c_idx)

                # Format total rows specially
                if is_total_row:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color=palette["total_text"])
                    cell.fill = PatternFill(start_color=palette["total_fill"], end_color=palette["total_fill"], fill_type="solid")
                    cell.border = total_bottom_border
                else:
                    cell.font = data_font
                    cell.fill = fill
                    cell.border = thin_border

                # Numeric auto-casting & Excel formula injection
                typed_val, align, num_format = cls._parse_cell_value(clean_val, c_idx, current_r, is_total_row, rows)
                cell.value = typed_val
                cell.alignment = align
                if num_format:
                    cell.number_format = num_format

            current_r += 1

        return current_r

    @classmethod
    def _parse_cell_value(
        cls,
        text: str,
        col_idx: int,
        current_row: int,
        is_total_row: bool,
        rows: Tuple[Tuple[str, ...], ...],
    ) -> Tuple[Any, Alignment, str]:
        """Convert string text into typed numbers, floats, currency, or Excel formulas."""
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        if not text:
            return "", align_left, ""

        # Auto-formula generation for total rows
        if is_total_row and (text.lower() in ("total", "sum", "") or text == "0"):
            if col_idx == 1:
                return text or "Total", align_left, ""
            col_letter = get_column_letter(col_idx)
            start_data_row = current_row - (len(rows) - 1)
            end_data_row = current_row - 1
            formula = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
            return formula, align_right, "#,##0.00"

        # Check for currency ($1,234.56 or €100)
        curr_match = re.match(r"^[\$€£]\s*([\d,]+(?:\.\d+)?)$", text)
        if curr_match:
            try:
                num = float(curr_match.group(1).replace(",", ""))
                symbol = text[0]
                return num, align_right, f"{symbol}#,##0.00"
            except ValueError:
                pass

        # Check for percentage (e.g. 15.5%)
        pct_match = re.match(r"^([\d.]+)\s*%$", text)
        if pct_match:
            try:
                num = float(pct_match.group(1)) / 100.0
                return num, align_right, "0.0%"
            except ValueError:
                pass

        # Check for integer
        if re.match(r"^-?\d+$", text):
            return int(text), align_right, "#,##0"

        # Check for float
        if re.match(r"^-?\d+\.\d+$", text):
            return float(text), align_right, "#,##0.00"

        return text, align_left, ""

    @classmethod
    def _auto_fit_columns(cls, ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
        """Autofit column widths with generous padding for high readability."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if cell.value and str(cell.value).startswith("="):
                    val_str = "123,456.78"  # Estimate for formula width
                max_len = max(max_len, len(val_str))

            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
