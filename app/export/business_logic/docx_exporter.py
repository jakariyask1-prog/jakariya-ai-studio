"""DOCX conversation exporter with layout rules isolated in templates."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.export.business_logic.asset_helpers import message_image_paths
from app.export.business_logic.base_exporter import BaseExporter
from app.export.business_logic.markdown_blocks import parse_markdown_blocks
from app.export.export_options import ExportOptions
from app.export.models import ExportMessage
from app.export.templates.docx_styles import add_title, configure_document, set_table_geometry


class DocxExporter(BaseExporter):
    """Render conversation content into a styled, editable Word document."""

    format_id = "docx"
    display_name = "Word document"
    file_extension = ".docx"

    def export(self, conversation, destination: Path, options: ExportOptions) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        configure_document(document)
        if not options.content_only:
            add_title(document, conversation.title)
        if options.include_metadata:
            for line in (
                f"Conversation ID: {conversation.id}",
                f"Created: {conversation.created_at}",
                f"Updated: {conversation.updated_at}",
            ):
                paragraph = document.add_paragraph(line)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = _rgb("6B7280")

        for message in conversation.messages:
            self._add_message(document, message, options)
        document.save(destination)
        return destination

    def _add_message(self, document, message: ExportMessage, options: ExportOptions) -> None:
        if not options.content_only:
            role_str = str(getattr(message.role, "value", message.role)).capitalize()
            document.add_heading(role_str, level=1)
        if options.include_timestamps:
            timestamp = document.add_paragraph(message.created_at)
            timestamp.paragraph_format.space_after = Pt(7)
            for run in timestamp.runs:
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = _rgb("6B7280")

        for block in parse_markdown_blocks(message.content):
            if block.kind == "code":
                paragraph = document.add_paragraph(block.text, style="Export Code")
                paragraph.paragraph_format.keep_together = True
            elif block.kind == "table":
                self._add_table(document, block.rows)
            else:
                self._add_paragraphs(document, block.text)

        if options.include_attachments:
            for attachment in message.attachments:
                if attachment.is_image and not options.include_images:
                    continue
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(f"Attachment: {attachment.original_name}").bold = True
                paragraph.add_run(f" ({attachment.media_type or attachment.extension or 'file'})")
        if options.include_images:
            for image_path in message_image_paths(message):
                try:
                    document.add_picture(str(image_path), width=Inches(5.6))
                    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except (OSError, ValueError):
                    document.add_paragraph(f"Image: {image_path.name}")

    @staticmethod
    def _add_paragraphs(document, text: str) -> None:
        for source_paragraph in filter(str.strip, text.split("\n\n")):
            source_paragraph = _remove_generated_image_markup(source_paragraph).strip()
            if not source_paragraph:
                continue
            heading_match = re.match(r"^(#{1,3})\s+(.+)$", source_paragraph)
            if heading_match and "\n" not in source_paragraph:
                document.add_heading(heading_match.group(2), level=len(heading_match.group(1)))
                continue
            paragraph = document.add_paragraph()
            _add_inline_markdown_runs(paragraph, source_paragraph)

    @staticmethod
    def _add_table(document, rows) -> None:
        column_count = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=column_count)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                cell = table.cell(row_index, column_index)
                cell.text = value
                if row_index == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        base_width, remainder = divmod(9360, column_count)
        widths = tuple(base_width + (1 if index < remainder else 0) for index in range(column_count))
        set_table_geometry(table, widths)


def _add_inline_markdown_runs(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|(?<!\*)\*[^*]+\*)", re.DOTALL)
    for part in pattern.split(text):
        if not part:
            continue
        run = paragraph.add_run(part.strip("*`"))
        if part.startswith("**"):
            run.bold = True
        elif part.startswith("`"):
            run.font.name = "Courier New"
        elif part.startswith("*"):
            run.italic = True


def _remove_generated_image_markup(text: str) -> str:
    return re.sub(r"<img[^>]*>|<br\s*/?>\s*<a[^>]*>.*?</a>", "", text, flags=re.IGNORECASE | re.DOTALL)


def _rgb(value: str):
    from docx.shared import RGBColor

    return RGBColor.from_string(value)
