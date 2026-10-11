import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_presentation_docx(output_path: Path):
    doc = docx.Document()

    # Set standard margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Color Palette
    PRIMARY = RGBColor(15, 23, 42)      # Deep Navy #0F172A
    SECONDARY = RGBColor(79, 70, 229)   # Indigo #4F46E5
    BODY_TEXT = RGBColor(51, 65, 85)     # Charcoal #334155
    MUTED = RGBColor(100, 116, 139)     # Slate #64748B
    PRIMARY_HEX = "0F172A"
    LIGHT_BG_HEX = "F8FAFC"
    ALT_ROW_HEX = "F1F5F9"
    BORDER_HEX = "CBD5E1"

    # Helper function for adding headings
    def add_custom_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        if level == 1:
            run.font.name = 'Segoe UI'
            run.font.size = Pt(18)
            run.font.color.rgb = PRIMARY
        elif level == 2:
            run.font.name = 'Segoe UI'
            run.font.size = Pt(14)
            run.font.color.rgb = SECONDARY
        elif level == 3:
            run.font.name = 'Segoe UI'
            run.font.size = Pt(12)
            run.font.color.rgb = PRIMARY
        return p

    def add_custom_paragraph(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.color.rgb = PRIMARY
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)
        r_text.font.color.rgb = BODY_TEXT
        return p

    def add_bullet_item(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = SECONDARY
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)
        r_text.font.color.rgb = BODY_TEXT
        return p

    # --- COVER BANNER / HEADER ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("JAKARIYA AI STUDIO")
    r_title.bold = True
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(28)
    r_title.font.color.rgb = PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    r_sub = sub_p.add_run("Enterprise Multi-Provider Desktop AI Ecosystem & Voice Intelligence Platform")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = SECONDARY

    meta_table = doc.add_table(rows=1, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    cell_data = [
        ("Prepared For:", "Client Presentation"),
        ("Version:", "v2.5 Enterprise"),
        ("Date:", "July 2026")
    ]
    for idx, (label, val) in enumerate(cell_data):
        cell = meta_table.rows[0].cells[idx]
        cell.width = Inches(2.1)
        set_cell_background(cell, LIGHT_BG_HEX)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label} ")
        r1.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = MUTED
        r2 = p.add_run(val)
        r2.bold = True
        r2.font.size = Pt(10)
        r2.font.color.rgb = PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- 1. EXECUTIVE OVERVIEW ---
    add_custom_heading("1. Executive Overview & Value Proposition", level=1)
    add_custom_paragraph(
        "Jakariya AI Studio is a state-of-the-art desktop artificial intelligence suite designed to unify world-class AI models, hands-free voice interaction, real-time web intelligence, and multi-format document generation into a single, high-performance desktop application.",
        bold_prefix=""
    )
    add_custom_paragraph(
        "Unlike standard web interfaces tied to a single AI vendor, Jakariya AI Studio provides an all-in-one ecosystem with zero vendor lock-in. It dynamically orchestrates 8 leading AI providers with built-in automated failover, local data security, multi-language speech recognition, and instant report generation.",
        bold_prefix="Core Advantage: "
    )

    # Key Highlights Grid
    add_bullet_item("Multi-Provider Orchestration:", "Seamless access to OpenAI GPT-4o, Claude 3.5 Sonnet, Gemini 2.5, DeepSeek R1, Llama 3.3 70B, and Mistral in one workspace.")
    add_bullet_item("Real-Time Voice Assistant (⚡ Live):", "Continuous hands-free conversation with active barge-in support and acoustic feedback suppression.")
    add_bullet_item("Autonomous Web Search:", "Integrated Tavily search engine retrieves real-time market data, current events, and live web snippets with clickable source citations.")
    add_bullet_item("Multi-Format Document Exporters:", "Instantly export chat sessions into publication-ready PDF, Word (.docx), Markdown (.md), HTML, and TXT files.")
    add_bullet_item("Data Privacy & Local Storage:", "Thread-safe SQLite database (WAL mode) ensures all conversations, attachments, and settings remain 100% private on your machine.")

    # --- 2. AI PROVIDER & MODEL CATALOG ---
    add_custom_heading("2. Multi-Provider AI Engine", level=1)
    add_custom_paragraph(
        "The application features a ProviderManager architecture that automatically monitors API health, rate limits, and authentication credentials. If a primary provider is unavailable, requests failover smoothly without interrupting user workflow."
    )

    # Provider Table
    p_table = doc.add_table(rows=9, cols=4)
    p_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Provider", "Key Models Included", "Speed / Performance", "Target Use Case"]
    hdr_cells = p_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], PRIMARY_HEX)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

    providers_data = [
        ("GitHub Models", "GPT-4o, GPT-4o-mini, Phi-4", "High Precision / Azure Infra", "Complex logic, deep reasoning & coding"),
        ("Groq Engine", "Llama 3.3 70B Versatile", "Ultra-Fast (300+ Tokens/sec)", "Real-time voice chat & rapid drafting"),
        ("Google AI Studio", "Gemini 2.5 Flash, Gemini 1.5 Pro", "Large Context Window (1M+)", "Document analysis & long-context research"),
        ("OpenRouter", "DeepSeek R1, DeepSeek V3", "Reasoning & Benchmark Leader", "Advanced math, logic & deep research"),
        ("Anthropic", "Claude 3.5 Sonnet, Claude 3 Opus", "Nuanced & Natural Prose", "Creative writing, coding & analysis"),
        ("Cerebras Engine", "Llama 3.3 70B", "Hardware Accelerated Speed", "Instant response generation"),
        ("Mistral AI", "Mistral Large, Codestral", "European Flagship AI", "Polyglot software development"),
        ("Kimi AI / Moonshot", "Kimi K1.5", "High Precision Context", "Multilingual chat & specialized tasks")
    ]

    col_widths = [Inches(1.5), Inches(1.8), Inches(1.6), Inches(1.6)]

    for row_idx, data in enumerate(providers_data, start=1):
        row_cells = p_table.rows[row_idx].cells
        bg_color = ALT_ROW_HEX if row_idx % 2 == 0 else LIGHT_BG_HEX
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            for run in p.runs:
                run.font.name = 'Calibri'
                run.font.size = Pt(9.5)
                run.font.color.rgb = BODY_TEXT
                if col_idx == 0:
                    run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- 3. VOICE & AUDIO PLATFORM ---
    add_custom_heading("3. Advanced Voice & Audio Intelligence", level=1)
    add_custom_paragraph(
        "Jakariya AI Studio turns desktop computing into an interactive voice assistant equipped with both single-turn dictation and continuous hands-free conversation."
    )

    add_custom_heading("Dual Voice Modes", level=2)
    add_bullet_item("Single-Turn Dictation Mode:", "Press the microphone button to dictate text. Supports auto-send or manual editing before submission.")
    add_bullet_item("Continuous Live Mode (⚡ Live):", "An active, hands-free conversation session. The assistant listens continuously and responds in real-time.")

    add_custom_heading("Barge-in & Feedback Suppression", level=2)
    add_custom_paragraph(
        "During live continuous speech, the system dynamically scales down microphone gain (0.15 attenuation during AI speech playback) to prevent acoustic feedback loop false barge-ins, while keeping user speech detection active so you can interrupt the assistant at any moment."
    )

    add_custom_heading("Multilingual Audio Support", level=2)
    add_custom_paragraph(
        "Native speech recognition (STT) and text-to-speech (TTS) engines support 11 global languages including Bangla (বাংলা), English, Spanish, French, German, Hindi, Japanese, Korean, Portuguese, Chinese, and Arabic."
    )

    # --- 4. EXPORT ENGINE & REAL-TIME WEB SEARCH ---
    add_custom_heading("4. Executive Productivity Features", level=1)

    add_custom_heading("Multi-Format Document Exporters", level=2)
    add_custom_paragraph(
        "Transform chat conversations into professional deliverables in one click. Exporters maintain message timestamps, visual formatting, embedded image attachments, and Bangla/Unicode characters:"
    )
    add_bullet_item("PDF Document (.pdf):", "Styled formal report layout with custom headers.")
    add_bullet_item("Microsoft Word (.docx):", "Editable Word document formatted for business sharing.")
    add_bullet_item("HTML Presentation (.html):", "Responsive web report with syntax highlighted code blocks.")
    add_bullet_item("Markdown & Text (.md, .txt):", "Clean text files for developer documentation.")
    add_bullet_item("One-Click Code Download:", "Download any generated code block directly to its proper extension (.py, .js, .cpp, etc.).")

    add_custom_heading("Autonomous Web Search Intelligence", level=2)
    add_custom_paragraph(
        "Integrated with Tavily API, the studio automatically detects queries requesting current news, market prices, weather forecasts, or live data. It injects live search context and web URL citations directly into the LLM context window."
    )

    # --- 5. ENTERPRISE ARCHITECTURE & QUALITY ASSURANCE ---
    add_custom_heading("5. Enterprise Architecture & Quality Guarantee", level=1)
    add_custom_paragraph(
        "Built strictly adhering to modern desktop software practices:"
    )

    add_bullet_item("PySide6 / Qt6 Native UI:", "High DPI responsive interface with custom fonts, glassmorphism aesthetics, and clean dark/light themes.")
    add_bullet_item("Thread-Safe SQLite Engine:", "ACID-compliant storage using WAL journal mode, busy timeouts, and automatic schema migrations.")
    add_bullet_item("Rigorous Quality Assurance:", "Covered by 128 automated Pytest unit tests achieving 100% pass rate across provider failovers, voice pipelines, UI headers, and exporter services.")

    # --- 6. COMMERCIAL PACKAGES & DEPLOYMENT ---
    add_custom_heading("6. Commercial Deployment & Customization Options", level=1)
    add_custom_paragraph(
        "Jakariya AI Studio is ready for white-label client deployment, custom branding, and enterprise integration:"
    )

    add_bullet_item("White-Label Branding:", "Customize logo, app branding, splash screens, and default color themes for your client organization.")
    add_bullet_item("Private Local LLM Integration:", "Optionally hook into local Ollama / LM Studio instances for 100% offline air-gapped corporate security.")
    add_bullet_item("Custom Workflow Integration:", "Add specialized exporters, custom prompt templates, or internal REST API hooks.")

    # --- FOOTER / CONCLUSION ---
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    callout_p = doc.add_paragraph()
    callout_p.paragraph_format.space_before = Pt(12)
    callout_p.paragraph_format.space_after = Pt(12)
    
    callout_table = doc.add_table(rows=1, cols=1)
    callout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = callout_table.rows[0].cells[0]
    c_cell.width = Inches(6.5)
    set_cell_background(c_cell, LIGHT_BG_HEX)
    set_cell_margins(c_cell, top=140, bottom=140, left=180, right=180)
    
    cp = c_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c1 = cp.add_run("Summary: ")
    r_c1.bold = True
    r_c1.font.name = 'Segoe UI'
    r_c1.font.size = Pt(11)
    r_c1.font.color.rgb = SECONDARY
    
    r_c2 = cp.add_run("Jakariya AI Studio combines 8 world-class AI providers, real-time voice & web intelligence, and multi-format document exporting into a polished, client-ready desktop application.")
    r_c2.font.name = 'Calibri'
    r_c2.font.size = Pt(10.5)
    r_c2.font.color.rgb = PRIMARY

    doc.save(str(output_path))
    print(f"Document successfully created at: {output_path}")

if __name__ == '__main__':
    out_file = Path("Jakariya_AI_Studio_Client_Presentation.docx").resolve()
    create_presentation_docx(out_file)
