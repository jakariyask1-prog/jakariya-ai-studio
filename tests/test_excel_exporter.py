import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

import openpyxl

from app.agent.tools.document_exporter_tool import DocumentExporterTool
from app.export.business_logic import ExcelExporter, build_default_exporter_registry
from app.export.export_options import ExportOptions
from app.export.models import ExportConversation, ExportMessage


class ExcelExporterTests(unittest.TestCase):
    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()
        self.home_patch = patch.object(Path, "home", classmethod(lambda cls: Path(self.temp_directory.name)))
        self.home_patch.start()
        (Path.home() / "Documents").mkdir()
        self.conversation = ExportConversation(
            id="chat-excel-1",
            title="Quarterly Sales Report",
            created_at="2026-08-07T10:00:00",
            updated_at="2026-08-07T10:05:00",
            messages=(
                ExportMessage(
                    id="msg-1",
                    role="assistant",
                    created_at="2026-08-07T10:01:00",
                    content=(
                        "# Financial Breakdown\n\n"
                        "| Region | Revenue | Units |\n"
                        "| --- | --- | --- |\n"
                        "| North | $15,000.00 | 120 |\n"
                        "| South | $22,500.50 | 180 |\n"
                        "| Total | $37,500.50 | 300 |"
                    ),
                ),
            ),
        )

    def tearDown(self):
        self.home_patch.stop()
        self.temp_directory.cleanup()

    def test_excel_export_creates_formatted_workbook(self):
        destination = Path(self.temp_directory.name) / "sales_report.xlsx"
        ExcelExporter().export(self.conversation, destination, ExportOptions("xlsx"))

        self.assertTrue(destination.exists())
        self.assertGreater(destination.stat().st_size, 1000)

        wb = openpyxl.load_workbook(str(destination))
        self.assertIn("Report", wb.sheetnames)
        ws = wb["Report"]

        # Check title block
        self.assertEqual(ws.cell(row=1, column=1).value, "Quarterly Sales Report")

        # Check headers
        self.assertEqual(ws.cell(row=5, column=1).value, "Region")
        self.assertEqual(ws.cell(row=5, column=2).value, "Revenue")
        self.assertEqual(ws.cell(row=5, column=3).value, "Units")

        # Check typed data values
        self.assertEqual(ws.cell(row=6, column=1).value, "North")
        self.assertEqual(ws.cell(row=6, column=2).value, 15000.0)
        self.assertEqual(ws.cell(row=6, column=3).value, 120)

    def test_excel_exporter_registry(self):
        registry = build_default_exporter_registry()
        exporter = registry.get("xlsx")
        self.assertIsInstance(exporter, ExcelExporter)
        self.assertEqual(exporter.file_extension, ".xlsx")

    def test_document_exporter_tool_xlsx(self):
        tool = DocumentExporterTool()
        self.assertIn(".xlsx", tool.metadata.supported_file_types)

        xlsx_path = "Documents/agent_report.xlsx"
        content = (
            "# Project Milestones\n\n"
            "| Task | Priority | Status |\n"
            "| --- | --- | --- |\n"
            "| Agent Engine | High | Completed |\n"
            "| Excel Export | Critical | Done |"
        )

        result = tool.execute(file_path=xlsx_path, content=content, title="Project Milestones")
        self.assertTrue(result.success, f"Tool failed: {result.error}")
        saved_path = Path(result.data["saved_path"])
        self.assertTrue(saved_path.exists())
        self.assertGreater(saved_path.stat().st_size, 1000)


    def test_document_exporter_tool_colorful_themes(self):
        tool = DocumentExporterTool()
        content = (
            "# Executive Color Roadmap 2026\n\n"
            "## Phase 1: Core Design System\n"
            "Deploy multi-color palettes and vibrant themes.\n\n"
            "> [!NOTE]\n"
            "> This is a colorful callout note.\n\n"
            "| Module | Color Palette | Status |\n"
            "| --- | --- | --- |\n"
            "| PDF Engine | Vibrant Purple | Active |\n"
            "| Word Engine | Emerald Green | Active |\n"
            "| Excel Engine | Sunset Orange | Active |"
        )

        for theme in ("vibrant", "colorful", "emerald", "sunset", "cyber"):
            pdf_path = f"Documents/report_{theme}.pdf"
            docx_path = f"Documents/report_{theme}.docx"
            html_path = f"Documents/report_{theme}.html"

            r_pdf = tool.execute(file_path=pdf_path, content=content, title=f"Report {theme}", theme=theme)
            self.assertTrue(r_pdf.success, f"PDF generation failed for theme '{theme}': {r_pdf.error}")

            r_docx = tool.execute(file_path=docx_path, content=content, title=f"Report {theme}", theme=theme)
            self.assertTrue(r_docx.success, f"DOCX generation failed for theme '{theme}': {r_docx.error}")

            r_html = tool.execute(file_path=html_path, content=content, title=f"Report {theme}", theme=theme)
            self.assertTrue(r_html.success, f"HTML generation failed for theme '{theme}': {r_html.error}")


if __name__ == "__main__":
    unittest.main()

