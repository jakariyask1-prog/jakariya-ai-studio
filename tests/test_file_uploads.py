import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from app.models.attachment import AttachmentInput
from app.uploads.file_service import FileService


class FileServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_directory.name)
        self.upload_directory = self.root / "uploads"
        self.service = FileService()
        self.path_patch = patch(
            "app.uploads.business_logic.file_storage.get_uploaded_files_dir",
            return_value=self.upload_directory,
        )
        self.path_patch.start()

    def tearDown(self):
        self.path_patch.stop()
        self.temp_directory.cleanup()

    def test_text_file_is_copied_and_added_to_context(self):
        source = self.root / "notes.txt"
        source.write_text("Project Falcon is scheduled for Friday.", encoding="utf-8")

        uploaded = self.service.add_file(source)

        self.assertTrue(uploaded.stored_path.exists())
        self.assertEqual(uploaded.original_name, "notes.txt")
        self.assertIn("Project Falcon", self.service.build_ai_context([uploaded]))
        self.assertEqual(uploaded.to_attachment_input().size_bytes, source.stat().st_size)

    def test_docx_text_and_table_content_are_extracted(self):
        source = self.root / "brief.docx"
        document = Document()
        document.add_paragraph("Release brief")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Owner"
        table.cell(0, 1).text = "Jakariya"
        document.save(source)

        uploaded = self.service.add_file(source)

        self.assertIn("Release brief", uploaded.text)
        self.assertIn("Owner | Jakariya", uploaded.text)

    def test_multiple_files_are_deduplicated_and_added(self):
        first = self.root / "first.txt"
        second = self.root / "second.md"
        first.write_text("one", encoding="utf-8")
        second.write_text("two", encoding="utf-8")

        uploaded = self.service.add_files([str(first), str(second), str(first)])

        self.assertEqual([file.original_name for file in uploaded], ["first.txt", "second.md"])
        self.assertTrue(all(file.stored_path.exists() for file in uploaded))

    def test_unsupported_extension_is_rejected(self):
        source = self.root / "program.exe"
        source.write_bytes(b"not allowed")

        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            self.service.add_file(source)

    def test_discard_files_removes_only_the_managed_upload_copy(self):
        source = self.root / "notes.txt"
        source.write_text("keep source", encoding="utf-8")
        uploaded = self.service.add_file(source)

        self.assertEqual(self.service.discard_files((uploaded,)), 1)
        self.assertFalse(uploaded.stored_path.exists())
        self.assertTrue(source.exists())

    def test_discard_files_never_deletes_a_path_outside_upload_storage(self):
        source = self.root / "notes.txt"
        source.write_text("keep source", encoding="utf-8")
        attachment = AttachmentInput("notes.txt", str(source), ".txt", source.stat().st_size)

        self.assertEqual(self.service.discard_files((attachment,)), 0)
        self.assertTrue(source.exists())

    def test_partial_batch_failure_removes_previously_copied_files(self):
        first = self.root / "first.txt"
        second = self.root / "second.txt"
        first.write_text("first", encoding="utf-8")
        second.write_text("second", encoding="utf-8")
        already_uploaded = self.service.add_file(first)

        with patch.object(
            self.service,
            "add_file",
            side_effect=[already_uploaded, ValueError("invalid second file")],
        ):
            with self.assertRaisesRegex(ValueError, "invalid second file"):
                self.service.add_files([str(first), str(second)])

        self.assertFalse(already_uploaded.stored_path.exists())

    def test_excel_and_code_file_upload_extractions(self):
        import openpyxl
        excel_source = self.root / "financials.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Q3 Summary"
        ws.cell(row=1, column=1, value="Metric")
        ws.cell(row=1, column=2, value="Amount")
        ws.cell(row=2, column=1, value="Revenue")
        ws.cell(row=2, column=2, value="$50,000")
        wb.save(str(excel_source))

        uploaded_excel = self.service.add_file(excel_source)
        self.assertIn("### Sheet: Q3 Summary", uploaded_excel.text)
        self.assertIn("Revenue | $50,000", uploaded_excel.text)

        py_source = self.root / "script.py"
        py_source.write_text("print('hello world')", encoding="utf-8")
        uploaded_py = self.service.add_file(py_source)
        self.assertEqual(uploaded_py.text, "print('hello world')")


if __name__ == "__main__":
    unittest.main()

