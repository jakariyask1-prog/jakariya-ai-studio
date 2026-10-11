"""End-to-end context coverage for DOCX, PDF, and Markdown agent attachments."""

from pathlib import Path
from unittest.mock import MagicMock

from docx import Document
from reportlab.pdfgen import canvas

from app.agent.agent_engine import AIAgentEngine
from app.providers.models.chat_response import ChatResponse
from app.uploads.file_service import FileService


def test_agent_receives_docx_pdf_and_markdown_attachment_text(tmp_path: Path, monkeypatch):
    upload_dir = tmp_path / "managed-uploads"
    monkeypatch.setattr(
        "app.uploads.business_logic.file_storage.get_uploaded_files_dir",
        lambda: upload_dir,
    )

    docx_path = tmp_path / "reference.docx"
    document = Document()
    document.add_heading("DOCX Reference Heading", level=1)
    document.add_paragraph("DOCX reference body text.")
    document.save(docx_path)

    pdf_path = tmp_path / "reference.pdf"
    pdf_canvas = canvas.Canvas(str(pdf_path))
    pdf_canvas.drawString(72, 720, "PDF reference body text.")
    pdf_canvas.save()

    markdown_path = tmp_path / "reference.md"
    markdown_path.write_text("# Markdown Reference Heading\n\nMarkdown reference body text.", encoding="utf-8")

    file_service = FileService()
    uploaded = file_service.add_files([str(docx_path), str(pdf_path), str(markdown_path)])
    prompt = file_service.build_message_content("Create a DOCX from these references.", uploaded)

    assert isinstance(prompt, str)
    assert "DOCX Reference Heading" in prompt
    assert "PDF reference body text." in prompt
    assert "Markdown Reference Heading" in prompt

    agent = AIAgentEngine()
    mock_provider = MagicMock()
    mock_provider.chat.return_value = ChatResponse(
        content="Final Answer: Reference files were received.",
        model="mock-model",
        provider_id="mock-provider",
    )
    agent.provider_manager = mock_provider

    agent.execute_task(prompt, model="mock-model")

    request = mock_provider.chat.call_args.args[0]
    assert request.messages[-1].content == prompt
