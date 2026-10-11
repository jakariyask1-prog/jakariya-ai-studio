import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.export.business_logic import DocxExporter, PdfExporter
from app.export.export_options import ExportOptions
from app.export.models import ExportConversation, ExportMessage


class DocumentExporterTests(unittest.TestCase):
    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()
        self.home_patch = patch.object(Path, "home", classmethod(lambda cls: Path(self.temp_directory.name)))
        self.home_patch.start()
        (Path.home() / "Documents").mkdir()
        self.conversation = ExportConversation(
            id="chat-1", title="Export document", created_at="2026-07-18T10:00:00",
            updated_at="2026-07-18T10:05:00",
            messages=(ExportMessage(
                id="message-1", role="assistant", created_at="2026-07-18T10:01:00",
                content=("## Result\n\nHere is **bold** and `inline code`.\n\n"
                         "```python\nprint('hello')\n```\n\n"
                         "| Name | Value |\n| --- | --- |\n| A | 1 |"),
            ),),
        )

    def tearDown(self):
        self.home_patch.stop()
        self.temp_directory.cleanup()

    def test_pdf_export_creates_a_readable_document(self):
        destination = Path(self.temp_directory.name) / "conversation.pdf"
        PdfExporter().export(self.conversation, destination, ExportOptions("pdf"))

        self.assertGreater(destination.stat().st_size, 500)
        reader = PdfReader(str(destination))
        self.assertEqual(len(reader.pages), 1)
        self.assertIn("Export document", reader.pages[0].extract_text())
        self.assertIn("print('hello')", reader.pages[0].extract_text())

    def test_docx_export_creates_a_styled_editable_document_with_a_table(self):
        destination = Path(self.temp_directory.name) / "conversation.docx"
        DocxExporter().export(self.conversation, destination, ExportOptions("docx"))

        document = Document(destination)
        paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Export document", paragraph_text)
        self.assertIn("print('hello')", paragraph_text)
        self.assertEqual(document.tables[0].cell(1, 0).text, "A")
        self.assertEqual(document.tables[0].cell(1, 1).text, "1")


    def test_document_exporter_tool_pdf_special_characters(self):
        from app.agent.tools.document_exporter_tool import DocumentExporterTool
        tool = DocumentExporterTool()
        pdf_path = "Documents/special.pdf"
        content_with_xml = "Test code: if x < 10 and y > 20 & z: print('<div>test</div>')"
        result = tool.execute(file_path=pdf_path, content=content_with_xml, title="Special <Char> Test")
        self.assertTrue(result.success, f"DocumentExporterTool failed: {result.error}")
        saved_path = Path(result.data["saved_path"])
        self.assertTrue(saved_path.exists())
        self.assertGreater(saved_path.stat().st_size, 500)

    def test_parse_tool_call_filename_normalization(self):
        from app.agent.agent_engine import AIAgentEngine
        engine = AIAgentEngine()
        text = 'Action: export_document\nAction Input: {"file_name": "report.pdf", "text": "sample data"}'
        parsed = engine._parse_tool_call(text)
        self.assertIsNotNone(parsed)
        tool_name, params = parsed
        self.assertEqual(tool_name, "export_document")
        self.assertEqual(params.get("file_path"), "report.pdf")
        self.assertEqual(params.get("content"), "sample data")


if __name__ == "__main__":
    unittest.main()
